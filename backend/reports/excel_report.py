import io
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelReportGenerator:
    """
    Autonomous Excel Workbook Generator using openpyxl.
    Builds structured, multi-sheet analytical workbooks with executive styling.
    """

    def generate(
        self,
        report_payload: Dict[str, Any],
        raw_df=None
    ) -> bytes:
        """Generate multi-tab Excel workbook bytes."""
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style definitions
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        subhead_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        subhead_font = Font(name="Calibri", size=11, bold=True, color="1E293B")
        regular_font = Font(name="Calibri", size=11, color="000000")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        def style_header_row(ws, row_idx, max_col):
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        def auto_fit_columns(ws):
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # ==========================================
        # SHEET 1: Summary
        # ==========================================
        ws_sum = wb.create_sheet(title="Executive Summary")
        ws_sum.views.sheetView[0].showGridLines = True
        ws_sum["A1"] = "AUTOANALYST AI - EXECUTIVE SUMMARY"
        ws_sum["A1"].font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
        
        meta = report_payload.get("metadata", {})
        ws_sum["A3"] = "Dataset:"
        ws_sum["B3"] = meta.get("dataset_name")
        ws_sum["A4"] = "Generated At:"
        ws_sum["B4"] = meta.get("generated_at")
        ws_sum["A5"] = "User Question:"
        ws_sum["B5"] = report_payload.get("4_user_question")
        for r in range(3, 6):
            ws_sum[f"A{r}"].font = Font(bold=True)

        ws_sum["A7"] = "Executive Overview:"
        ws_sum["A7"].font = Font(bold=True, size=12)
        ws_sum["A8"] = report_payload.get("1_executive_summary", "")

        ws_sum["A10"] = "Key Strategic Recommendations:"
        ws_sum["A10"].font = Font(bold=True, size=12)
        recs = report_payload.get("14_recommendations", [])
        for i, rec in enumerate(recs, start=11):
            ws_sum[f"A{i}"] = f"{i-10}. {rec}"

        auto_fit_columns(ws_sum)

        # ==========================================
        # SHEET 2: Dataset Info
        # ==========================================
        ws_info = wb.create_sheet(title="Dataset Info")
        ws_info.views.sheetView[0].showGridLines = True
        ws_info.append(["Property", "Value"])
        style_header_row(ws_info, 1, 2)

        ds = report_payload.get("2_dataset_overview", {})
        ws_info.append(["Total Rows", ds.get("rows", 0)])
        ws_info.append(["Total Columns", ds.get("columns", 0)])
        ws_info.append(["Memory Usage (MB)", ds.get("memory_mb", 0.0)])
        ws_info.append(["Numerical Columns", ", ".join(ds.get("numerical_columns", []))])
        ws_info.append(["Categorical Columns", ", ".join(ds.get("categorical_columns", []))])
        ws_info.append(["Datetime Columns", ", ".join(ds.get("datetime_columns", []))])
        auto_fit_columns(ws_info)

        # ==========================================
        # SHEET 3: Data Quality
        # ==========================================
        ws_dq = wb.create_sheet(title="Data Quality")
        ws_dq.views.sheetView[0].showGridLines = True
        ws_dq.append(["Quality Metric", "Measurement"])
        style_header_row(ws_dq, 1, 2)

        dq = report_payload.get("3_data_quality", {})
        ws_dq.append(["Overall Quality Score", f"{dq.get('quality_score', 100)}/100"])
        ws_dq.append(["Quality Rating", dq.get("rating", "Good")])
        ws_dq.append(["Total Missing Cells", dq.get("total_missing_cells", 0)])
        ws_dq.append(["Missing Data Percentage", f"{dq.get('missing_percentage', 0.0)}%"])
        ws_dq.append(["Duplicate Rows", dq.get("duplicate_rows", 0)])
        ws_dq.append(["Outliers Flagged", dq.get("outliers_count", 0)])
        auto_fit_columns(ws_dq)

        # ==========================================
        # SHEET 4: KPIs
        # ==========================================
        ws_kpi = wb.create_sheet(title="Key KPIs")
        ws_kpi.views.sheetView[0].showGridLines = True
        ws_kpi.append(["KPI Dimension", "Metric / Attribute", "Calculated Value", "Context / Sub-Metric"])
        style_header_row(ws_kpi, 1, 4)

        kpis = report_payload.get("7_key_kpis", {})
        pm = kpis.get("primary_metric")
        if pm:
            ws_kpi.append(["Primary Focal Metric", f"Total {pm['name']}", pm['total'], f"Mean: {pm['mean']} | Median: {pm['median']}"])
            ws_kpi.append(["Primary Focal Metric", f"Max {pm['name']}", pm['max'], f"Min: {pm['min']} | Std: {pm['std']}"])

        top = kpis.get("top_performer")
        if top:
            ws_kpi.append(["Segmentation", f"Top {top['dimension']}", str(top['category']), f"Total: {top['total']} | Mean: {top['mean']}"])

        low = kpis.get("lowest_performer")
        if low:
            ws_kpi.append(["Segmentation", f"Lowest {low['dimension']}", str(low['category']), f"Total: {low['total']} | Mean: {low['mean']}"])

        pg = kpis.get("period_growth")
        if pg:
            ws_kpi.append(["Growth Trajectory", "Overall Growth", f"{pg['overall_growth_pct']}%", f"Recent Period: {pg['recent_period_growth_pct']}%"])

        anom = kpis.get("anomalies")
        if anom:
            ws_kpi.append(["Anomaly Detection", "Anomalies Count", anom['count'], f"{anom['rate_pct']}% of data"])

        auto_fit_columns(ws_kpi)

        # ==========================================
        # SHEET 5: Important Findings
        # ==========================================
        ws_find = wb.create_sheet(title="Important Findings")
        ws_find.views.sheetView[0].showGridLines = True
        ws_find.append(["Category", "Finding Statement"])
        style_header_row(ws_find, 1, 2)

        for fact in report_payload.get("8_important_findings", []):
            ws_find.append(["Observed Fact", fact])
        for note in report_payload.get("9_statistical_analysis", {}).get("statistical_notes", []):
            ws_find.append(["Statistical Result", note])
        for interp in report_payload.get("13_business_insights", []):
            ws_find.append(["AI Interpretation", interp])
        auto_fit_columns(ws_find)

        # ==========================================
        # SHEET 6: Anomalies / Predictions
        # ==========================================
        anom_data = report_payload.get("12_anomalies", {}).get("top_records", [])
        if anom_data:
            ws_anom = wb.create_sheet(title="Anomalies")
            ws_anom.views.sheetView[0].showGridLines = True
            headers = list(anom_data[0].keys())
            ws_anom.append(headers)
            style_header_row(ws_anom, 1, len(headers))
            for row in anom_data:
                ws_anom.append([row.get(h) for h in headers])
            auto_fit_columns(ws_anom)

        # ==========================================
        # SHEET 7: Dataset Sample / Slice
        # ==========================================
        if raw_df is not None and len(raw_df) > 0:
            ws_data = wb.create_sheet(title="Dataset Sample")
            ws_data.views.sheetView[0].showGridLines = True
            cols = list(raw_df.columns)
            ws_data.append(cols)
            style_header_row(ws_data, 1, len(cols))
            for _, row in raw_df.head(200).iterrows():
                ws_data.append([str(v) if not isinstance(v, (int, float)) else v for v in row])
            auto_fit_columns(ws_data)

        buffer = io.BytesIO()
        wb.save(buffer)
        excel_bytes = buffer.getvalue()
        buffer.close()
        return excel_bytes
