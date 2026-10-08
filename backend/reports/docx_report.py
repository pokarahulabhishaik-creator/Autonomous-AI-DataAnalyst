import io
from typing import Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

class DocxReportGenerator:
    """
    Autonomous DOCX Report Generator using python-docx.
    Builds styled Microsoft Word documents with executive formatting and KPI tables.
    """

    def generate(self, report_payload: Dict[str, Any]) -> bytes:
        """Compile report dictionary into formatted DOCX bytes."""
        doc = Document()

        # Adjust margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Title
        meta = report_payload.get("metadata", {})
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("AUTOANALYST AI")
        title_run.font.size = Pt(24)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy

        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run("Autonomous Data Analysis & Business Intelligence Report")
        sub_run.font.size = Pt(12)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 116, 139)

        meta_p = doc.add_paragraph()
        meta_run = meta_p.add_run(f"Dataset: {meta.get('dataset_name')}  |  Generated: {meta.get('generated_at')}")
        meta_run.font.size = Pt(9.5)
        meta_run.font.color.rgb = RGBColor(71, 85, 105)

        # Question Callout
        question = report_payload.get("4_user_question", "")
        q_p = doc.add_paragraph()
        q_p.paragraph_format.left_indent = Inches(0.3)
        q_head = q_p.add_run("INVESTIGATION OBJECTIVE:\n")
        q_head.font.bold = True
        q_head.font.size = Pt(10)
        q_head.font.color.rgb = RGBColor(37, 99, 235)
        q_text = q_p.add_run(f"\"{question}\"")
        q_text.font.italic = True
        q_text.font.size = Pt(10.5)

        # 1. Executive Summary
        h1 = doc.add_heading("1. Executive Summary", level=2)
        h1.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        doc.add_paragraph(report_payload.get("1_executive_summary", ""))

        # 2. Key Computed KPIs Table
        h2 = doc.add_heading("2. Key Computed KPIs (Empirical Results)", level=2)
        h2.runs[0].font.color.rgb = RGBColor(30, 58, 138)

        kpis = report_payload.get("7_key_kpis", {})
        kpi_rows = [["Indicator", "Calculated Value", "Benchmark / Context"]]

        pm = kpis.get("primary_metric")
        if pm:
            kpi_rows.append([f"Total {pm['name']}", f"{pm['total']:,}", f"Mean: {pm['mean']:,} | Max: {pm['max']:,}"])

        top = kpis.get("top_performer")
        if top:
            kpi_rows.append([f"Top {top['dimension']}", str(top['category']), f"Total: {top['total']:,} (Mean: {top['mean']:,})"])

        low = kpis.get("lowest_performer")
        if low:
            kpi_rows.append([f"Lowest {low['dimension']}", str(low['category']), f"Total: {low['total']:,} (Mean: {low['mean']:,})"])

        pg = kpis.get("period_growth")
        if pg:
            kpi_rows.append(["Period Growth", f"{pg['overall_growth_pct']}%", f"Recent: {pg['recent_period_growth_pct']}%"])

        anom = kpis.get("anomalies")
        if anom:
            kpi_rows.append(["Anomalous Records", f"{anom['count']}", f"{anom['rate_pct']}% of transactions"])

        table = doc.add_table(rows=len(kpi_rows), cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Light Shading Accent 1'

        for r_idx, row_data in enumerate(kpi_rows):
            for c_idx, val in enumerate(row_data):
                cell = table.cell(r_idx, c_idx)
                cell.text = str(val)
                if r_idx == 0:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True

        doc.add_paragraph()

        # 3. Important Observed Findings
        h3 = doc.add_heading("3. Important Observed Findings", level=2)
        h3.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        for fact in report_payload.get("8_important_findings", []):
            p = doc.add_paragraph(style='List Bullet')
            r_bold = p.add_run("Observed Fact: ")
            r_bold.font.bold = True
            p.add_run(fact)

        # 4. Statistical Analysis
        h4 = doc.add_heading("4. Statistical Analysis & Hypothesis Testing", level=2)
        h4.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        for note in report_payload.get("9_statistical_analysis", {}).get("statistical_notes", []):
            p = doc.add_paragraph(style='List Bullet')
            r_bold = p.add_run("Statistical Finding: ")
            r_bold.font.bold = True
            p.add_run(note)

        # 5. AI Business Interpretations
        h5 = doc.add_heading("5. Analytical Interpretations", level=2)
        h5.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        for interp in report_payload.get("13_business_insights", []):
            p = doc.add_paragraph(style='List Bullet')
            r_bold = p.add_run("AI Interpretation: ")
            r_bold.font.bold = True
            p.add_run(interp)

        # 6. Actionable Recommendations
        h6 = doc.add_heading("6. Strategic Business Recommendations", level=2)
        h6.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        for rec in report_payload.get("14_recommendations", []):
            p = doc.add_paragraph(style='List Bullet')
            r_bold = p.add_run("Action Item: ")
            r_bold.font.bold = True
            p.add_run(rec)

        # 7. Dataset Overview & Data Quality Audit
        h7 = doc.add_heading("7. Dataset Quality & Audit Overview", level=2)
        h7.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        dq = report_payload.get("3_data_quality", {})
        ds = report_payload.get("2_dataset_overview", {})
        dq_p = doc.add_paragraph()
        dq_p.add_run(
            f"Dataset dimensions: {ds.get('rows', 0)} rows, {ds.get('columns', 0)} columns. "
            f"Overall Data Quality Score: {dq.get('quality_score', 100)}/100 ({dq.get('rating')}). "
            f"Missing cells: {dq.get('total_missing_cells', 0)} ({dq.get('missing_percentage', 0)}%), "
            f"Duplicates: {dq.get('duplicate_rows', 0)}, Outliers: {dq.get('outliers_count', 0)}."
        )

        # 8. Conclusion
        h8 = doc.add_heading("8. Conclusion", level=2)
        h8.runs[0].font.color.rgb = RGBColor(30, 58, 138)
        doc.add_paragraph(report_payload.get("15_conclusion", ""))

        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()
        return docx_bytes
